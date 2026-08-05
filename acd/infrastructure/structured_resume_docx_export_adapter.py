"""python-docx implementation of the structured resume export port."""

from __future__ import annotations

from pathlib import Path
from tempfile import NamedTemporaryFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt

from acd.application.effective_structured_resume_docx_export import (
    StructuredResumeDocxExportPort,
)
from acd.application.structured_resume_snapshot import StructuredResumeSnapshot


class PythonDocxStructuredResumeExportAdapter(StructuredResumeDocxExportPort):
    """Render a validated snapshot as an ATS-friendly single-column DOCX."""

    def export(self, snapshot: StructuredResumeSnapshot, destination_path: Path) -> None:
        """Write through a sibling temporary file before atomically replacing output."""
        document = Document()
        self._configure(document)
        self._render(document, snapshot)
        temporary_path: Path | None = None
        try:
            with NamedTemporaryFile(
                suffix=".docx", prefix=f".{destination_path.stem}.", dir=destination_path.parent, delete=False
            ) as temporary:
                temporary_path = Path(temporary.name)
            document.save(temporary_path)
            if not temporary_path.is_file() or temporary_path.stat().st_size == 0:
                raise OSError("DOCX export did not produce a valid file.")
            temporary_path.replace(destination_path)
        finally:
            if temporary_path is not None and temporary_path.exists():
                temporary_path.unlink()

    @staticmethod
    def _configure(document: Document) -> None:
        section = document.sections[0]
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        styles = document.styles
        normal = styles["Normal"]
        normal.font.name = "Aptos"
        normal.font.size = Pt(10)
        heading = styles["Heading 1"]
        heading.font.name = "Aptos Display"
        heading.font.size = Pt(12)
        heading.font.bold = True
        heading.paragraph_format.space_before = Pt(10)
        heading.paragraph_format.space_after = Pt(3)
        if "Resume Name" not in styles:
            name_style = styles.add_style("Resume Name", WD_STYLE_TYPE.PARAGRAPH)
            name_style.font.name = "Aptos Display"
            name_style.font.size = Pt(20)
            name_style.font.bold = True
        document.core_properties.title = "Currículo"
        document.core_properties.subject = "Currículo estruturado"

    def _render(self, document: Document, snapshot: StructuredResumeSnapshot) -> None:
        self._paragraph(document, snapshot.identity.full_name, "Resume Name")
        self._paragraph(document, snapshot.identity.professional_title)
        contact = self._join(
            snapshot.identity.location,
            snapshot.contact.phone,
            snapshot.contact.email,
            snapshot.contact.linkedin or snapshot.contact.portfolio or snapshot.contact.website,
        )
        self._paragraph(document, contact)
        self._section(document, "RESUMO PROFISSIONAL", (snapshot.summary,))
        self._section(document, "COMPETÊNCIAS", snapshot.skills, bullets=True)
        if snapshot.experiences:
            document.add_heading("EXPERIÊNCIA PROFISSIONAL", level=1)
            for experience in snapshot.experiences:
                self._paragraph(document, self._join(experience.role, experience.company))
                self._paragraph(document, self._join(experience.location, experience.start_date, experience.end_date))
                self._paragraph(document, experience.summary)
                self._section(document, None, experience.achievements, bullets=True)
        self._section(document, "FORMAÇÃO ACADÊMICA", tuple(
            self._join(item.degree, item.field, item.institution, item.start_date, item.end_date, item.status, item.details)
            for item in snapshot.education
        ), bullets=True)
        self._section(document, "CERTIFICAÇÕES", tuple(
            self._join(item.name, item.issuer, item.issued_date, item.credential_id, item.url)
            for item in snapshot.certifications
        ), bullets=True)
        self._section(document, "CURSOS", tuple(
            self._join(item.name, item.provider, item.completed_date, item.details)
            for item in snapshot.courses
        ), bullets=True)
        self._section(document, "IDIOMAS", tuple(
            self._join(item.name, item.proficiency) for item in snapshot.languages
        ), bullets=True)
        if snapshot.projects:
            document.add_heading("PROJETOS", level=1)
            for project in snapshot.projects:
                self._paragraph(document, self._join(project.name, project.url))
                self._paragraph(document, project.description)
                self._section(document, None, project.highlights, bullets=True)
        for section in sorted(snapshot.additional_sections, key=lambda item: item.order):
            self._section(document, section.title, (*section.paragraphs, *section.items), bullets=bool(section.items))

    @staticmethod
    def _section(document: Document, title: str | None, values: tuple[str | None, ...], *, bullets: bool = False) -> None:
        content = tuple(value for value in values if value)
        if not content:
            return
        if title:
            document.add_heading(title, level=1)
        for value in content:
            document.add_paragraph(value, style="List Bullet" if bullets else None)

    @staticmethod
    def _paragraph(document: Document, value: str | None, style: str | None = None) -> None:
        if value:
            document.add_paragraph(value, style=style)

    @staticmethod
    def _join(*values: str | None) -> str:
        return " | ".join(value for value in values if value)
