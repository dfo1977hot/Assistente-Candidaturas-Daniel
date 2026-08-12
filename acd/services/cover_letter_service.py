from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
from typing import Any

from acd.domain.entities.cover_letter_version import CoverLetterVersion
from acd.infrastructure.repositories.cover_letter_repository import CoverLetterRepository
from acd.services.curriculum_service import CurriculumService
from acd.services.job_service import JobService
from acd.services.settings_service import SettingsService


@dataclass(frozen=True, slots=True)
class CoverLetterContext:
    job_id: int
    curriculum_id: int
    company: str
    title: str
    recruiter: str
    recruiter_email: str
    job_notes: str
    curriculum_name: str
    curriculum_version: str
    curriculum_content: str


class CoverLetterService:
    """Regras de negócio, versionamento, IA e exportação de cartas."""

    LETTER_TYPES = (
        "Carta de Apresentação",
        "Carta de Motivação",
        "E-mail de Candidatura",
        "Mensagem para Recrutador/LinkedIn",
        "Carta Espontânea",
    )
    LANGUAGES = ("Português", "Inglês", "Espanhol")
    TONES = ("Profissional", "Executivo", "Técnico", "Objetivo", "Persuasivo")
    LENGTHS = ("Curta", "Média", "Completa")
    STATUSES = ("Rascunho", "Gerada", "Revisada", "Aprovada", "Enviada", "Arquivada")

    def __init__(
        self,
        repository: CoverLetterRepository | None = None,
        job_service: JobService | None = None,
        curriculum_service: CurriculumService | None = None,
        settings_service: SettingsService | None = None,
        client: Any | None = None,
    ) -> None:
        self.repository = repository or CoverLetterRepository()
        self.job_service = job_service or JobService()
        self.curriculum_service = curriculum_service or CurriculumService()
        self.settings_service = settings_service or SettingsService()
        self._client = client

    def list_letters(self) -> list[CoverLetterVersion]:
        return self.repository.get_all()

    def search_letters(self, query: str) -> list[CoverLetterVersion]:
        return self.repository.search(query)

    def get_letter(self, letter_id: int) -> CoverLetterVersion | None:
        return self.repository.get_by_id(letter_id)

    def save_letter(
        self,
        *,
        letter_id: int | None,
        job_id: int,
        curriculum_id: int,
        version: str,
        letter_type: str,
        language: str,
        tone: str,
        length: str,
        subject: str,
        status: str,
        content: str,
        notes: str,
    ) -> CoverLetterVersion:
        self._validate(job_id=job_id, curriculum_id=curriculum_id, content=content)
        if letter_id is None:
            entity = CoverLetterVersion(
                job_id=job_id,
                application_id=self.repository.latest_application_id_for_job(job_id),
                curriculum_id=curriculum_id,
                version=version.strip() or self.next_version(
                    job_id=job_id,
                    curriculum_id=curriculum_id,
                ),
                letter_type=letter_type,
                language=language,
                tone=tone,
                length=length,
                subject=subject.strip(),
                status=status,
                content=content.strip(),
                notes=notes.strip(),
            )
            return self.repository.create(entity)

        entity = self.repository.get_by_id(letter_id)
        if entity is None:
            raise ValueError("Carta não encontrada.")
        entity.job_id = job_id
        entity.application_id = (
            entity.application_id
            or self.repository.latest_application_id_for_job(job_id)
        )
        entity.curriculum_id = curriculum_id
        entity.version = version.strip() or entity.version
        entity.letter_type = letter_type
        entity.language = language
        entity.tone = tone
        entity.length = length
        entity.subject = subject.strip()
        entity.status = status
        entity.content = content.strip()
        entity.notes = notes.strip()
        return self.repository.update(entity)

    def delete_letter(self, letter_id: int) -> bool:
        return self.repository.delete(letter_id)

    def next_version(self, *, job_id: int, curriculum_id: int) -> str:
        versions = self.repository.versions_for_context(
            job_id=job_id,
            curriculum_id=curriculum_id,
        )
        highest = (1, -1)
        for item in versions:
            match = re.fullmatch(
                r"[vV]?(?P<major>\d+)\.(?P<minor>\d+)",
                item.version.strip(),
            )
            if match:
                pair = (int(match.group("major")), int(match.group("minor")))
                if pair > highest:
                    highest = pair
        if highest == (1, -1):
            return "V1.0"
        return f"V{highest[0]}.{highest[1] + 1}"

    def build_context(self, *, job_id: int, curriculum_id: int) -> CoverLetterContext:
        job = self.job_service.get_job(job_id)
        curriculum = self.curriculum_service.repository.get_by_id(curriculum_id)
        if job is None:
            raise ValueError("Selecione uma vaga válida.")
        if curriculum is None:
            raise ValueError("Selecione um currículo válido.")

        curriculum_content = (
            curriculum.structured_content_json
            or curriculum.description
            or ""
        )
        return CoverLetterContext(
            job_id=job.id,
            curriculum_id=curriculum.id,
            company=getattr(job.company, "name", "") or "",
            title=job.title,
            recruiter=job.recruiter or "",
            recruiter_email=job.recruiter_email or "",
            job_notes=job.notes or "",
            curriculum_name=curriculum.name,
            curriculum_version=curriculum.version,
            curriculum_content=curriculum_content,
        )

    def generate_letter(
        self,
        *,
        job_id: int,
        curriculum_id: int,
        letter_type: str,
        language: str,
        tone: str,
        length: str,
    ) -> CoverLetterVersion:
        context = self.build_context(job_id=job_id, curriculum_id=curriculum_id)
        client = self._client or self._create_client()
        prompt = self._build_prompt(
            context=context,
            letter_type=letter_type,
            language=language,
            tone=tone,
            length=length,
        )
        response = client.responses.create(
            model="gpt-5-mini",
            input=prompt,
        )
        content = str(getattr(response, "output_text", "") or "").strip()
        if not content:
            raise RuntimeError("A IA não retornou conteúdo para a carta.")

        version = self.next_version(job_id=job_id, curriculum_id=curriculum_id)
        subject = self._default_subject(
            context=context,
            letter_type=letter_type,
        )
        entity = CoverLetterVersion(
            job_id=job_id,
            application_id=self.repository.latest_application_id_for_job(job_id),
            curriculum_id=curriculum_id,
            version=version,
            letter_type=letter_type,
            language=language,
            tone=tone,
            length=length,
            subject=subject,
            status="Gerada",
            content=content,
            notes="",
            explanation=(
                "Gerada com base na vaga, empresa, recrutador e currículo selecionado."
            ),
        )
        return self.repository.create(entity)

    def export_docx(self, letter_id: int, destination: str | Path) -> Path:
        letter = self._require_letter(letter_id)
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError(
                "A biblioteca python-docx não está instalada."
            ) from exc

        path = Path(destination)
        document = Document()
        if letter.subject:
            document.add_heading(letter.subject, level=1)
        for paragraph in letter.content.split("\n"):
            document.add_paragraph(paragraph)
        document.save(path)
        return path

    def export_pdf(self, letter_id: int, destination: str | Path) -> Path:
        letter = self._require_letter(letter_id)
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        except ImportError as exc:
            raise RuntimeError("A biblioteca reportlab não está instalada.") from exc

        path = Path(destination)
        styles = getSampleStyleSheet()
        story = []
        if letter.subject:
            story.extend([Paragraph(letter.subject, styles["Title"]), Spacer(1, 12)])
        for paragraph in letter.content.split("\n"):
            safe = (
                paragraph.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )
            story.extend([Paragraph(safe or " ", styles["BodyText"]), Spacer(1, 8)])
        SimpleDocTemplate(str(path), pagesize=A4).build(story)
        return path

    def _create_client(self) -> Any:
        try:
            from openai import OpenAI
        except ImportError as exc:
            raise RuntimeError(
                "A biblioteca openai não está instalada no ambiente."
            ) from exc
        api_key = self.settings_service.get_api_key("openai")
        if not api_key:
            raise RuntimeError(
                "Configure a chave da OpenAI na página Configurações."
            )
        return OpenAI(api_key=api_key)

    @staticmethod
    def _build_prompt(
        *,
        context: CoverLetterContext,
        letter_type: str,
        language: str,
        tone: str,
        length: str,
    ) -> str:
        return f"""
Você é um especialista em recrutamento e comunicação profissional.
Escreva uma {letter_type} altamente personalizada.

Idioma: {language}
Tom: {tone}
Tamanho: {length}

Vaga:
Empresa: {context.company}
Cargo: {context.title}
Recrutador: {context.recruiter}
E-mail do recrutador: {context.recruiter_email}
Descrição/observações da vaga:
{context.job_notes}

Currículo selecionado:
Nome: {context.curriculum_name}
Versão: {context.curriculum_version}
Conteúdo disponível:
{context.curriculum_content}

Regras:
- Não invente experiências, resultados, competências, nomes ou números.
- Use apenas informações existentes no contexto.
- Priorize resultados mensuráveis e competências aderentes à vaga.
- Evite aberturas genéricas e clichês.
- Não escreva explicações antes ou depois da carta.
- Retorne somente o texto final da carta.
""".strip()

    @staticmethod
    def _default_subject(
        *,
        context: CoverLetterContext,
        letter_type: str,
    ) -> str:
        if "E-mail" in letter_type or "Mensagem" in letter_type:
            return f"Candidatura — {context.title} — {context.company}".strip(" —")
        return f"{context.title} — {context.company}".strip(" —")

    @staticmethod
    def _validate(*, job_id: int, curriculum_id: int, content: str) -> None:
        if job_id <= 0:
            raise ValueError("Selecione uma vaga.")
        if curriculum_id <= 0:
            raise ValueError("Selecione um currículo.")
        if not content.strip():
            raise ValueError("O conteúdo da carta não pode ficar vazio.")

    def _require_letter(self, letter_id: int) -> CoverLetterVersion:
        letter = self.repository.get_by_id(letter_id)
        if letter is None:
            raise ValueError("Carta não encontrada.")
        return letter
