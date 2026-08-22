from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
import hashlib
import os
from pathlib import Path
import re
import shutil
from uuid import uuid4
import zipfile

from docx import Document
from docx.text.paragraph import Paragraph

from acd.core.logger import logger
from acd.domain.entities.curriculum import Curriculum
from acd.infrastructure.repositories.curriculum_repository import CurriculumRepository


class CurriculumDocumentError(ValueError):
    """Erro funcional ao manipular documentos de currículo."""


@dataclass(frozen=True)
class DocumentIntegrity:
    status: str
    message: str


class CurriculumDocumentService:
    ALLOWED_EXTENSIONS = frozenset({".docx", ".pdf"})
    MAX_FILE_SIZE_BYTES = 20 * 1024 * 1024
    _VERSION_SUFFIX_PATTERN = re.compile(
        r"_v\d+(?:\.\d+)?$",
        flags=re.IGNORECASE,
    )
    _BLOCK_MARKER_PATTERN = re.compile(
        r"^\s*(?:[-*]\s*)?(?:\*\*)?"
        r"\[\[\s*BLOCO\s+(\d+)\s*\]\](?:\*\*)?\s*:?[ \t]*(.*)$",
        flags=re.IGNORECASE,
    )

    def __init__(
        self,
        repository: CurriculumRepository | None = None,
        project_root: Path | None = None,
    ) -> None:
        self.repository = repository or CurriculumRepository()
        self.project_root = (project_root or Path.cwd()).resolve()
        self.storage_root = (
            self.project_root / "data" / "documents" / "curricula"
        ).resolve()
        self.storage_root.mkdir(parents=True, exist_ok=True)

    def attach_document(
        self,
        curriculum_id: int,
        source_path: str | Path,
        *,
        original_name: str | None = None,
    ) -> Curriculum:
        curriculum = self._require_curriculum(curriculum_id)
        source = Path(source_path).expanduser().resolve()
        self._validate_source(source)
        destination = self._new_destination(source)
        shutil.copy2(source, destination)
        previous = self.resolve_path(curriculum.file_relative_path)

        try:
            self._apply_metadata(
                curriculum,
                source,
                destination,
                original_name=original_name,
            )
            updated = self.repository.update(curriculum)
        except Exception:
            destination.unlink(missing_ok=True)
            raise

        if previous and previous != destination:
            previous.unlink(missing_ok=True)

        logger.info("Documento anexado ao currículo %s", curriculum_id)
        return updated

    def remove_document(self, curriculum_id: int) -> Curriculum:
        curriculum = self._require_curriculum(curriculum_id)
        document = self.resolve_path(curriculum.file_relative_path)
        self._clear_metadata(curriculum)
        updated = self.repository.update(curriculum)

        if document:
            document.unlink(missing_ok=True)

        return updated

    def duplicate_document(
        self,
        source: Curriculum,
        target: Curriculum,
        *,
        original_name: str | None = None,
    ) -> Curriculum:
        source_path = self.resolve_path(source.file_relative_path)
        if source_path is None or not source_path.exists():
            return target

        display_name = (
            original_name
            or source.file_original_name
            or source_path.name
        )
        return self.attach_document(
            target.id,
            source_path,
            original_name=display_name,
        )

    def materialize_optimized_document(
        self,
        source: Curriculum,
        target: Curriculum,
        *,
        optimized_content: str,
        version: str,
    ) -> Curriculum:
        """Cria DOCX otimizado sobre uma cópia do documento original."""
        source_path = self.resolve_path(source.file_relative_path)
        if source_path is None or not source_path.exists():
            return target

        if source_path.suffix.lower() != ".docx":
            raise CurriculumDocumentError(
                "A otimização com preservação integral da formatação "
                "é suportada apenas para currículos DOCX."
            )

        content = optimized_content.strip()
        if not content:
            raise CurriculumDocumentError(
                "A otimização não retornou conteúdo para o currículo."
            )

        original_name = (
            source.file_original_name
            or source_path.name
        )
        versioned_name = self.versioned_original_name(
            original_name,
            version,
        )

        original_checksum = self.calculate_sha256(source_path)
        destination = self._new_destination(Path(versioned_name))
        previous = self.resolve_path(target.file_relative_path)

        shutil.copy2(source_path, destination)

        try:
            self._apply_optimized_text_to_docx(
                destination,
                content,
            )

            if self.calculate_sha256(source_path) != original_checksum:
                raise CurriculumDocumentError(
                    "O documento original foi alterado durante a otimização."
                )

            self._apply_metadata(
                target,
                source_path,
                destination,
                original_name=versioned_name,
            )
            updated = self.repository.update(target)
        except Exception:
            destination.unlink(missing_ok=True)
            raise

        if previous and previous != destination:
            previous.unlink(missing_ok=True)

        logger.info(
            "Documento otimizado materializado: origem=%s destino=%s versão=%s",
            source.id,
            target.id,
            version,
        )
        return updated

    def extract_text_blocks(self, curriculum_id: int) -> tuple[str, ...]:
        """Extrai blocos editáveis de um DOCX mantendo sua ordem lógica."""
        curriculum = self._require_curriculum(curriculum_id)
        path = self.resolve_path(curriculum.file_relative_path)

        if path is None or not path.exists():
            return ()

        if path.suffix.lower() != ".docx":
            return ()

        try:
            document = Document(path)
        except (OSError, ValueError) as exc:
            raise CurriculumDocumentError(
                "Não foi possível ler o DOCX do currículo."
            ) from exc

        return tuple(
            paragraph.text.strip()
            for paragraph in self._editable_paragraphs(document)
        )

    def open_document(self, curriculum_id: int) -> None:
        curriculum = self._require_curriculum(curriculum_id)
        path = self.resolve_path(curriculum.file_relative_path)

        if path is None or not path.exists():
            raise CurriculumDocumentError("Arquivo não encontrado.")

        try:
            os.startfile(path)  # type: ignore[attr-defined]
        except OSError as exc:
            raise CurriculumDocumentError(
                "Não foi possível abrir o arquivo no aplicativo padrão do Windows."
            ) from exc

    def verify_integrity(self, curriculum: Curriculum) -> DocumentIntegrity:
        if not curriculum.file_relative_path:
            return DocumentIntegrity("none", "Sem arquivo")

        path = self.resolve_path(curriculum.file_relative_path)
        if path is None or not path.exists():
            return DocumentIntegrity("missing", "Arquivo não encontrado")

        try:
            checksum = self.calculate_sha256(path)
        except OSError:
            return DocumentIntegrity("error", "Erro de leitura")

        if curriculum.file_sha256 and checksum != curriculum.file_sha256:
            return DocumentIntegrity("changed", "Arquivo alterado externamente")

        return DocumentIntegrity("verified", "Verificada")

    def resolve_path(self, relative_path: str | None) -> Path | None:
        if not relative_path:
            return None

        candidate = (self.project_root / relative_path).resolve()
        self._assert_managed(candidate)
        return candidate

    @staticmethod
    def calculate_sha256(path: Path) -> str:
        digest = hashlib.sha256()

        with path.open("rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(chunk)

        return digest.hexdigest()

    @classmethod
    def versioned_original_name(
        cls,
        original_name: str,
        version: str,
    ) -> str:
        """Mantém o nome-base original e acrescenta somente a versão atual."""
        original = Path(original_name)
        suffix = original.suffix
        stem = original.stem if suffix else original.name
        stem = cls._VERSION_SUFFIX_PATTERN.sub("", stem)

        normalized_version = version.strip()
        if not normalized_version:
            raise CurriculumDocumentError("Versão do currículo inválida.")

        if not normalized_version.lower().startswith("v"):
            normalized_version = f"v{normalized_version}"

        return f"{stem}_{normalized_version}{suffix}"

    def _require_curriculum(self, curriculum_id: int) -> Curriculum:
        curriculum = self.repository.get_by_id(curriculum_id)
        if curriculum is None:
            raise CurriculumDocumentError("Currículo não encontrado.")
        return curriculum

    def _validate_source(self, source: Path) -> None:
        if not source.is_file():
            raise CurriculumDocumentError("Selecione um arquivo existente.")

        extension = source.suffix.lower()

        if extension not in self.ALLOWED_EXTENSIONS:
            raise CurriculumDocumentError(
                "Formato não permitido. Selecione um arquivo DOCX ou PDF."
            )

        if source.stat().st_size > self.MAX_FILE_SIZE_BYTES:
            raise CurriculumDocumentError(
                "O arquivo excede o limite máximo permitido de 20 MB."
            )

        if extension == ".pdf":
            with source.open("rb") as stream:
                if stream.read(5) != b"%PDF-":
                    raise CurriculumDocumentError(
                        "O arquivo selecionado não é um PDF válido."
                    )
            return

        try:
            with zipfile.ZipFile(source) as archive:
                names = set(archive.namelist())
                valid = (
                    "[Content_Types].xml" in names
                    and any(name.startswith("word/") for name in names)
                )
                if not valid:
                    raise CurriculumDocumentError(
                        "O arquivo selecionado não é um DOCX válido."
                    )
        except zipfile.BadZipFile as exc:
            raise CurriculumDocumentError(
                "O arquivo selecionado não é um DOCX válido."
            ) from exc

    def _new_destination(self, source: Path) -> Path:
        stem = re.sub(
            r"[^a-zA-Z0-9_-]+",
            "-",
            source.stem,
        ).strip("-").lower()

        destination = self.storage_root / (
            f"{uuid4()}-{stem[:80] or 'curriculo'}"
            f"{source.suffix.lower()}"
        )
        self._assert_managed(destination)
        return destination

    def _apply_metadata(
        self,
        curriculum: Curriculum,
        source: Path,
        destination: Path,
        *,
        original_name: str | None = None,
    ) -> None:
        extension = source.suffix.lower()

        curriculum.file_original_name = original_name or source.name
        curriculum.file_relative_path = destination.relative_to(
            self.project_root
        ).as_posix()
        curriculum.file_extension = extension
        curriculum.file_mime_type = {
            ".pdf": "application/pdf",
            ".docx": (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        }[extension]
        curriculum.file_size_bytes = destination.stat().st_size
        curriculum.file_sha256 = self.calculate_sha256(destination)
        curriculum.file_attached_at = datetime.now(UTC).replace(tzinfo=None)

    @staticmethod
    def _clear_metadata(curriculum: Curriculum) -> None:
        for attribute in (
            "file_original_name",
            "file_relative_path",
            "file_extension",
            "file_mime_type",
            "file_size_bytes",
            "file_sha256",
            "file_attached_at",
        ):
            setattr(curriculum, attribute, None)

    def _assert_managed(self, path: Path) -> None:
        try:
            path.resolve().relative_to(self.storage_root)
        except ValueError as exc:
            raise CurriculumDocumentError(
                "O caminho do documento está fora da pasta administrada."
            ) from exc

    def _apply_optimized_text_to_docx(
        self,
        path: Path,
        optimized_content: str,
    ) -> None:
        try:
            document = Document(path)
        except (OSError, ValueError) as exc:
            raise CurriculumDocumentError(
                "Não foi possível abrir a cópia do DOCX para otimização."
            ) from exc

        paragraphs = self._editable_paragraphs(document)

        if not paragraphs:
            raise CurriculumDocumentError(
                "O DOCX original não possui blocos textuais editáveis."
            )

        original_blocks = tuple(paragraph.text for paragraph in paragraphs)
        optimized_blocks = self._reconcile_optimized_blocks(
            optimized_content,
            original_blocks,
        )

        for paragraph, optimized_text in zip(
            paragraphs,
            optimized_blocks,
            strict=True,
        ):
            self._replace_paragraph_text_preserving_runs(
                paragraph,
                optimized_text,
            )

        try:
            document.save(path)
        except OSError as exc:
            raise CurriculumDocumentError(
                "Não foi possível salvar o DOCX otimizado."
            ) from exc

    @classmethod
    def _reconcile_optimized_blocks(
        cls,
        content: str,
        original_blocks: tuple[str, ...],
    ) -> tuple[str, ...]:
        """Reconcilia a resposta da IA com a estrutura fixa do DOCX.

        Quando existem marcadores, a numeração é a fonte de verdade: blocos
        podem chegar fora de ordem e blocos omitidos são preservados com o
        texto original. Duplicidades, índices inválidos e conteúdo estrutural
        ambíguo continuam sendo rejeitados.
        """
        if not original_blocks:
            return ()

        indexed = cls._parse_indexed_optimized_blocks(
            content,
            expected_count=len(original_blocks),
        )
        if indexed is not None:
            return tuple(
                indexed.get(index, original_blocks[index - 1])
                for index in range(1, len(original_blocks) + 1)
            )

        legacy_blocks = cls._parse_optimized_blocks(content)
        if len(legacy_blocks) != len(original_blocks):
            raise CurriculumDocumentError(
                "A IA alterou a quantidade de blocos do currículo. "
                "A versão otimizada não foi criada para evitar perda "
                "da formatação original."
            )
        return legacy_blocks

    @classmethod
    def _parse_indexed_optimized_blocks(
        cls,
        content: str,
        *,
        expected_count: int,
    ) -> dict[int, str] | None:
        """Lê blocos numerados sem depender da ordem em que a IA os retornou."""
        lines = [
            line.rstrip()
            for line in content.splitlines()
            if line.strip()
        ]
        if not lines:
            return None

        if not any(cls._BLOCK_MARKER_PATTERN.match(line) for line in lines):
            return None

        blocks: dict[int, str] = {}
        current_index: int | None = None
        current_parts: list[str] = []

        def commit_current() -> None:
            nonlocal current_index, current_parts
            if current_index is None:
                return
            if current_index in blocks:
                raise CurriculumDocumentError(
                    "A IA retornou um bloco duplicado do currículo."
                )
            blocks[current_index] = " ".join(current_parts).strip()
            current_index = None
            current_parts = []

        for line in lines:
            match = cls._BLOCK_MARKER_PATTERN.match(line)
            if match is not None:
                commit_current()
                index = int(match.group(1))
                if index < 1 or index > expected_count:
                    raise CurriculumDocumentError(
                        "A IA retornou um número de bloco que não existe "
                        "no currículo original."
                    )
                current_index = index
                current_parts = [match.group(2).strip()]
                continue

            stripped = line.strip()
            if current_index is None:
                if stripped in {"```", "```text", "```plaintext"}:
                    continue
                raise CurriculumDocumentError(
                    "A IA retornou conteúdo fora dos blocos esperados."
                )
            if stripped == "```":
                continue
            current_parts.append(stripped)

        commit_current()
        return blocks

    @classmethod
    def _parse_optimized_blocks(
        cls,
        content: str,
    ) -> tuple[str, ...]:
        """Compatibilidade com respostas legadas sem marcadores de bloco."""
        lines = [
            line.rstrip()
            for line in content.splitlines()
            if line.strip()
        ]
        return tuple(line.strip() for line in lines)

    @classmethod
    def _editable_paragraphs(
        cls,
        document: object,
    ) -> tuple[Paragraph, ...]:
        paragraphs: list[Paragraph] = []

        for paragraph in getattr(document, "paragraphs", ()):
            if paragraph.text.strip():
                paragraphs.append(paragraph)

        for table in getattr(document, "tables", ()):
            cls._append_table_paragraphs(table, paragraphs)

        return tuple(paragraphs)

    @classmethod
    def _append_table_paragraphs(
        cls,
        table: object,
        destination: list[Paragraph],
    ) -> None:
        seen_cells: set[int] = set()

        for row in getattr(table, "rows", ()):
            for cell in getattr(row, "cells", ()):
                cell_key = id(getattr(cell, "_tc", cell))
                if cell_key in seen_cells:
                    continue

                seen_cells.add(cell_key)

                for paragraph in getattr(cell, "paragraphs", ()):
                    if paragraph.text.strip():
                        destination.append(paragraph)

                for nested_table in getattr(cell, "tables", ()):
                    cls._append_table_paragraphs(
                        nested_table,
                        destination,
                    )

    @staticmethod
    def _replace_paragraph_text_preserving_runs(
        paragraph: Paragraph,
        new_text: str,
    ) -> None:
        """Troca texto mantendo os objetos Run e suas propriedades visuais."""
        runs = list(paragraph.runs)

        if not runs:
            paragraph.add_run(new_text)
            return

        original_lengths = [
            max(len(run.text), 1)
            for run in runs
        ]
        total_original = sum(original_lengths)

        consumed = 0
        cumulative_weight = 0

        for index, (run, weight) in enumerate(
            zip(runs, original_lengths, strict=True)
        ):
            cumulative_weight += weight

            if index == len(runs) - 1:
                run.text = new_text[consumed:]
                continue

            boundary = round(
                len(new_text) * cumulative_weight / total_original
            )
            boundary = max(consumed, min(boundary, len(new_text)))
            run.text = new_text[consumed:boundary]
            consumed = boundary