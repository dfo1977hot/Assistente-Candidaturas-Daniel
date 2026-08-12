from __future__ import annotations

from collections.abc import Callable, Sequence
import csv
import os
from pathlib import Path
import re
from typing import Protocol

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from acd.services.analytics_service import AnalyticsRecord


class CancellationGateway(Protocol):
    @property
    def is_cancellation_requested(self) -> bool: ...


ProgressCallback = Callable[[object], None]


class ExportCancelled(RuntimeError):
    """Signal a cooperative export cancellation after partial cleanup."""


class AnalyticsExportService:
    """Export an already-filtered AnalyticsService projection without database access."""

    HEADERS = (
        "Empresa", "Cargo", "Plataforma/origem", "Status da candidatura",
        "Data de criação", "Última atualização", "Remuneração oferecida",
        "Remuneração ideal", "Fit score", "Currículo associado",
        "Carta associada", "Workflow/status operacional",
    )

    def safe_filename(self, label: str, extension: str) -> str:
        normalized = re.sub(r"[^A-Za-z0-9_-]+", "-", label.strip()).strip("-")
        return f"acd-analytics-{normalized or 'dados'}.{extension.lstrip('.')}"

    def export_csv(
        self,
        records: Sequence[AnalyticsRecord],
        destination: str | Path,
        *,
        progress: ProgressCallback | None = None,
        cancellation: CancellationGateway | None = None,
    ) -> Path:
        target, partial = self._paths(destination)
        try:
            with partial.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(self.HEADERS)
                for index, record in enumerate(records, start=1):
                    self._check_cancelled(cancellation)
                    writer.writerow(self._row(record))
                    self._progress(progress, index, len(records))
            self._check_cancelled(cancellation)
            os.replace(partial, target)
            return target
        except Exception:
            partial.unlink(missing_ok=True)
            raise

    def export_xlsx(
        self,
        records: Sequence[AnalyticsRecord],
        destination: str | Path,
        *,
        progress: ProgressCallback | None = None,
        cancellation: CancellationGateway | None = None,
    ) -> Path:
        target, partial = self._paths(destination)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Analytics"
        sheet.append(self.HEADERS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(self.HEADERS))}1"
        try:
            for index, record in enumerate(records, start=1):
                self._check_cancelled(cancellation)
                sheet.append(self._row(record))
                row_number = index + 1
                for column in (5, 6):
                    sheet.cell(row_number, column).number_format = "dd/mm/yyyy hh:mm"
                for column in (7, 8):
                    sheet.cell(row_number, column).number_format = 'R$ #,##0.00'
                self._progress(progress, index, len(records))
            for index, header in enumerate(self.HEADERS, start=1):
                sheet.column_dimensions[get_column_letter(index)].width = min(
                    32, max(14, len(header) + 2)
                )
            self._check_cancelled(cancellation)
            workbook.save(partial)
            self._check_cancelled(cancellation)
            os.replace(partial, target)
            return target
        except Exception:
            partial.unlink(missing_ok=True)
            raise
        finally:
            workbook.close()

    @staticmethod
    def _row(record: AnalyticsRecord) -> tuple[object, ...]:
        return (
            record.company, record.title, record.source, record.application_status,
            record.created_at, record.updated_at, record.salary_offered,
            record.salary_ideal, record.fit_score, record.curriculum,
            record.cover_letter, record.workflow_status,
        )

    @staticmethod
    def _paths(destination: str | Path) -> tuple[Path, Path]:
        target = Path(destination).expanduser().resolve()
        target.parent.mkdir(parents=True, exist_ok=True)
        return target, target.with_name(f".{target.name}.partial")

    @staticmethod
    def _check_cancelled(cancellation: CancellationGateway | None) -> None:
        if cancellation is not None and cancellation.is_cancellation_requested:
            raise ExportCancelled("Exportação cancelada pelo usuário.")

    @staticmethod
    def _progress(callback: ProgressCallback | None, current: int, total: int) -> None:
        if callback is not None:
            callback({"current": current, "total": total, "percent": int(current / max(1, total) * 100)})
