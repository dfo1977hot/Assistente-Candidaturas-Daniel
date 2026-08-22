from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

from acd.domain.entities.curriculum import Curriculum
from acd.services.curriculum_document_service import (
    CurriculumDocumentService,
)
from acd.services.curriculum_service import CurriculumService


class _Repository:
    def __init__(self, *curricula: Curriculum) -> None:
        self.items = {
            curriculum.id: curriculum
            for curriculum in curricula
        }

    def get_by_id(
        self,
        curriculum_id: int,
    ) -> Curriculum | None:
        return self.items.get(curriculum_id)

    def update(
        self,
        curriculum: Curriculum,
    ) -> Curriculum:
        self.items[curriculum.id] = curriculum
        return curriculum


def _curriculum(
    curriculum_id: int,
    *,
    name: str = "Daniel Freitas",
    version: str = "v1.0",
) -> Curriculum:
    curriculum = Curriculum(
        name=name,
        version=version,
        language="pt-BR",
        description="",
    )
    curriculum.id = curriculum_id
    return curriculum


def _create_source_docx(path: Path) -> None:
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.81)
    section.left_margin = Inches(0.63)
    section.right_margin = Inches(0.67)

    normal = document.styles["Normal"]
    normal.font.name = "Century Gothic"
    normal.font.size = Pt(10)

    name = document.add_paragraph()
    name.paragraph_format.space_after = Pt(7)
    run = name.add_run("Daniel Freitas Oliveira")
    run.font.name = "Century Gothic"
    run.font.size = Pt(14)
    run.bold = True

    summary = document.add_paragraph()
    summary.paragraph_format.space_before = Pt(3)
    summary.paragraph_format.space_after = Pt(5)
    summary.paragraph_format.line_spacing = 1.15
    run = summary.add_run(
        "Engenheiro de Produção com experiência em logística."
    )
    run.font.name = "Century Gothic"
    run.font.size = Pt(10)

    document.save(path)


def test_versioned_original_name_preserves_base_name() -> None:
    assert (
        CurriculumDocumentService.versioned_original_name(
            "Daniel_Freitas_Oliveira.docx",
            "v2.0",
        )
        == "Daniel_Freitas_Oliveira_v2.0.docx"
    )

    assert (
        CurriculumDocumentService.versioned_original_name(
            "Daniel_Freitas_Oliveira_v2.0.docx",
            "v3.0",
        )
        == "Daniel_Freitas_Oliveira_v3.0.docx"
    )


def test_optimized_curriculum_versions_advance_major_version() -> None:
    assert CurriculumService._next_optimized_version(
        "v1.0"
    ) == "v2.0"

    assert CurriculumService._next_optimized_version(
        "v2.0"
    ) == "v3.0"

    assert CurriculumService._next_optimized_version(
        "v7.4"
    ) == "v8.0"


def test_extract_text_blocks_reads_original_docx(
    tmp_path: Path,
) -> None:
    source = _curriculum(1)
    repository = _Repository(source)

    service = CurriculumDocumentService(
        repository,  # type: ignore[arg-type]
        project_root=tmp_path,
    )

    source_path = (
        tmp_path
        / "Daniel_Freitas_Oliveira.docx"
    )
    _create_source_docx(source_path)

    service.attach_document(
        source.id,
        source_path,
    )

    blocks = service.extract_text_blocks(
        source.id
    )

    assert blocks == (
        "Daniel Freitas Oliveira",
        "Engenheiro de Produção com experiência em logística.",
    )


def test_materialized_optimized_docx_preserves_layout_and_original(
    tmp_path: Path,
) -> None:
    source = _curriculum(1)

    target = _curriculum(
        2,
        version="v2.0",
    )

    repository = _Repository(
        source,
        target,
    )

    service = CurriculumDocumentService(
        repository,  # type: ignore[arg-type]
        project_root=tmp_path,
    )

    source_path = (
        tmp_path
        / "Daniel_Freitas_Oliveira_Engenheiro_Producao.docx"
    )

    _create_source_docx(source_path)

    source = service.attach_document(
        source.id,
        source_path,
    )

    repository.items[source.id] = source

    managed_source = service.resolve_path(
        source.file_relative_path
    )

    assert managed_source is not None

    checksum_before = service.calculate_sha256(
        managed_source
    )

    optimized = service.materialize_optimized_document(
        source,
        target,
        optimized_content=(
            "[[BLOCO 1]] Daniel Freitas Oliveira\n"
            "[[BLOCO 2]] Engenheiro de Produção com experiência "
            "em logística, melhoria contínua e gestão de transportes."
        ),
        version="v2.0",
    )

    checksum_after = service.calculate_sha256(
        managed_source
    )

    assert checksum_after == checksum_before

    assert (
        optimized.file_original_name
        == (
            "Daniel_Freitas_Oliveira_"
            "Engenheiro_Producao_v2.0.docx"
        )
    )

    optimized_path = service.resolve_path(
        optimized.file_relative_path
    )

    assert optimized_path is not None
    assert optimized_path != managed_source
    assert optimized_path.exists()

    original = Document(managed_source)
    generated = Document(optimized_path)

    assert original.paragraphs[1].text == (
        "Engenheiro de Produção com experiência em logística."
    )

    assert generated.paragraphs[1].text == (
        "Engenheiro de Produção com experiência em logística, "
        "melhoria contínua e gestão de transportes."
    )

    original_section = original.sections[0]
    generated_section = generated.sections[0]

    assert (
        generated_section.top_margin
        == original_section.top_margin
    )

    assert (
        generated_section.bottom_margin
        == original_section.bottom_margin
    )

    assert (
        generated_section.left_margin
        == original_section.left_margin
    )

    assert (
        generated_section.right_margin
        == original_section.right_margin
    )

    assert (
        generated.paragraphs[0].paragraph_format.space_after
        == original.paragraphs[0].paragraph_format.space_after
    )

    assert (
        generated.paragraphs[1].paragraph_format.space_before
        == original.paragraphs[1].paragraph_format.space_before
    )

    assert (
        generated.paragraphs[1].paragraph_format.space_after
        == original.paragraphs[1].paragraph_format.space_after
    )

    assert (
        generated.paragraphs[0].runs[0].font.name
        == "Century Gothic"
    )

    assert (
        generated.paragraphs[0].runs[0].font.size
        == Pt(14)
    )

    assert generated.paragraphs[0].runs[0].bold is True

    assert (
        generated.paragraphs[1].runs[0].font.name
        == "Century Gothic"
    )

    assert (
        generated.paragraphs[1].runs[0].font.size
        == Pt(10)
    )


def test_materialization_preserves_missing_indexed_blocks(
    tmp_path: Path,
) -> None:
    source = _curriculum(1)

    target = _curriculum(
        2,
        version="v2.0",
    )

    repository = _Repository(
        source,
        target,
    )

    service = CurriculumDocumentService(
        repository,  # type: ignore[arg-type]
        project_root=tmp_path,
    )

    source_path = tmp_path / "Daniel.docx"
    _create_source_docx(source_path)

    source = service.attach_document(
        source.id,
        source_path,
    )

    repository.items[source.id] = source

    updated = service.materialize_optimized_document(
        source,
        target,
        optimized_content=(
            "[[BLOCO 1]] Apenas um bloco otimizado"
        ),
        version="v2.0",
    )

    optimized_path = service.resolve_path(
        updated.file_relative_path
    )

    assert optimized_path is not None
    assert optimized_path.exists()

    blocks = service.extract_text_blocks(updated.id)

    assert blocks[0] == "Apenas um bloco otimizado"
    assert len(blocks) > 1

def test_suggested_export_filename_uses_attached_versioned_name(
    tmp_path: Path,
) -> None:
    curriculum = _curriculum(
        1,
        name="Daniel Freitas Oliveira",
        version="v2.0",
    )

    curriculum.file_original_name = (
        "Daniel_Freitas_Oliveira_"
        "Engenheiro_Producao_v2.0.docx"
    )

    repository = _Repository(curriculum)

    service = CurriculumService(
        repository=repository,  # type: ignore[arg-type]
        prompt_repository=object(),
    )

    service.document_service = CurriculumDocumentService(
        repository,  # type: ignore[arg-type]
        project_root=tmp_path,
    )

    assert service.suggested_export_filename(
        curriculum.id
    ) == (
        "Daniel_Freitas_Oliveira_"
        "Engenheiro_Producao_v2.0.docx"
    )


def test_suggested_export_filename_preserves_original_v1_name(
    tmp_path: Path,
) -> None:
    curriculum = _curriculum(
        1,
        name="Daniel Freitas Oliveira",
        version="v1.0",
    )

    curriculum.file_original_name = (
        "Daniel_Freitas_Oliveira_"
        "Engenheiro_Producao.docx"
    )

    repository = _Repository(curriculum)

    service = CurriculumService(
        repository=repository,  # type: ignore[arg-type]
        prompt_repository=object(),
    )

    service.document_service = CurriculumDocumentService(
        repository,  # type: ignore[arg-type]
        project_root=tmp_path,
    )

    assert service.suggested_export_filename(
        curriculum.id
    ) == (
        "Daniel_Freitas_Oliveira_"
        "Engenheiro_Producao.docx"
    )


def test_suggested_export_filename_builds_version_when_document_name_missing(
    tmp_path: Path,
) -> None:
    curriculum = _curriculum(
        1,
        name="Daniel_Freitas_Oliveira",
        version="v3.0",
    )

    repository = _Repository(curriculum)

    service = CurriculumService(
        repository=repository,  # type: ignore[arg-type]
        prompt_repository=object(),
    )

    service.document_service = CurriculumDocumentService(
        repository,  # type: ignore[arg-type]
        project_root=tmp_path,
    )

    assert (
        service.suggested_export_filename(
            curriculum.id
        )
        == "Daniel_Freitas_Oliveira_v3.0.docx"
    )
