from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook
import pytest

from acd.services.analytics_export_service import AnalyticsExportService, ExportCancelled
from acd.services.analytics_service import AnalyticsFilters, AnalyticsRecord, AnalyticsService
from tests.services.test_dashboard_analytics_1_0 import OperationalRepository, real_data


def enriched_data(now: datetime):
    data = real_data(now)
    for job in data["jobs"]:
        job["salary_min"] = 8000
        job["salary_max"] = 10000
    for application in data["applications"]:
        application["curriculum_id"] = 1
        application["cover_letter_id"] = None
    data["curricula"] = [{"id": 1, "name": "Currículo Principal"}]
    data["ats_scores"] = [
        {"id": 1, "application_id": 1, "total_score": 82.5, "calculated_at": now}
    ]
    data["cover_letters"][0].update({"application_id": 1, "job_id": 1})
    for execution in data["workflow_executions"]:
        execution["application_id"] = 1 if execution["id"] == 3 else None
    return data


def record() -> AnalyticsRecord:
    now = datetime(2026, 8, 12, 12, tzinfo=UTC).replace(tzinfo=None)
    return AnalyticsRecord(
        application_id=1,
        job_id=2,
        company_id=3,
        company="ACME",
        title="Pessoa Desenvolvedora",
        source="LinkedIn",
        application_status="Aplicada",
        created_at=now,
        updated_at=now,
        salary_offered=8000.0,
        salary_ideal=10000.0,
        fit_score=82.5,
        curriculum="Currículo Principal",
        cover_letter="Carta #4",
        workflow_status="Aguardando usuário",
    )


def test_kpi_and_funnel_drilldowns_reuse_the_filtered_projection():
    now = datetime(2026, 8, 12, 12, tzinfo=UTC)
    service = AnalyticsService(OperationalRepository(enriched_data(now)))
    snapshot = service.get_dashboard(
        AnalyticsFilters(period="Todo o período", company_id=1, source="LinkedIn"), now=now
    )

    assert len(snapshot.records) == snapshot.kpis["applications_total"] == 3
    assert len(snapshot.drilldowns["applications_total"]) == 3
    assert all(item.company == "ACME" for item in snapshot.drilldowns["applications_total"])
    funnel = snapshot.drilldowns["funnel:Entrevista RH"]
    assert len(funnel) == 1
    assert funnel[0].status == "Entrevista RH"
    assert snapshot.records[0].salary_offered == 8000
    assert snapshot.records[0].salary_ideal == 10000
    assert snapshot.records[0].fit_score == 82.5
    assert not hasattr(snapshot.records[0], "notes")


def test_csv_is_utf8_bom_and_contains_only_supplied_filtered_records(tmp_path):
    destination = tmp_path / "analytics.csv"
    service = AnalyticsExportService()

    service.export_csv((record(),), destination)

    raw = destination.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    text = raw.decode("utf-8-sig")
    assert "ACME" in text
    assert text.count("Pessoa Desenvolvedora") == 1


def test_xlsx_has_headers_filter_freeze_dimensions_and_native_types(tmp_path):
    destination = tmp_path / "analytics.xlsx"
    service = AnalyticsExportService()

    service.export_xlsx((record(),), destination)
    workbook = load_workbook(destination)
    sheet = workbook.active

    assert tuple(cell.value for cell in sheet[1]) == service.HEADERS
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:L1"
    assert isinstance(sheet.cell(2, 5).value, datetime)
    assert sheet.cell(2, 7).value == 8000.0
    assert sheet.cell(2, 8).value == 10000.0
    workbook.close()


class Cancellation:
    def __init__(self):
        self.calls = 0

    @property
    def is_cancellation_requested(self):
        self.calls += 1
        return self.calls >= 2


@pytest.mark.parametrize(("method", "suffix"), [("export_csv", ".csv"), ("export_xlsx", ".xlsx")])
def test_cancellation_removes_partial_and_never_replaces_destination(tmp_path, method, suffix):
    destination = tmp_path / f"analytics{suffix}"
    destination.write_bytes(b"existing")
    service = AnalyticsExportService()

    with pytest.raises(ExportCancelled):
        getattr(service, method)((record(), record()), destination, cancellation=Cancellation())

    assert destination.read_bytes() == b"existing"
    assert not destination.with_name(f".{destination.name}.partial").exists()


def test_export_io_error_is_propagated_without_database_side_effect(tmp_path, monkeypatch):
    destination = tmp_path / "analytics.csv"
    service = AnalyticsExportService()
    monkeypatch.setattr(Path, "open", lambda *args, **kwargs: (_ for _ in ()).throw(OSError("disk")))

    with pytest.raises(OSError, match="disk"):
        service.export_csv((record(),), destination)

    assert not hasattr(service, "repository")
